from __future__ import annotations

import hashlib
import importlib.util
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "codex-skill"
TEMPLATE = SKILL / "assets/return-template.docx"


class TemplateReferenceTests(unittest.TestCase):
    def test_template_and_reference_cover_every_nonempty_cell(self) -> None:
        self.assertEqual(TEMPLATE.read_bytes(), (ROOT / "★倫理申請_修正対応依頼.docx").read_bytes())
        reference = (SKILL / "references/rules/r5-office-standard-text.md").read_text(encoding="utf-8")
        self.assertIn(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), reference)
        document = Document(TEMPLATE)
        self.assertEqual([(len(t.rows), len(t.columns)) for t in document.tables], [(52, 2), (1, 1)])
        body = reference.split("## 定型文（項目順）", 1)[1].split("## B・C枠等の定型文", 1)[0]
        sections = {}
        for section in re.split(r"(?m)^### ", body)[1:]:
            heading, content = section.split("\n", 1)
            quoted = "\n".join(line[2:] if line.startswith("> ") else ""
                               for line in content.splitlines() if line.startswith(">"))
            sections[heading] = quoted.strip()
        filled = [row for row in document.tables[0].rows if row.cells[1].text.strip()]
        self.assertEqual(len(sections), len(filled))
        for row in filled:
            heading = re.sub(r"^(\d+)\.\s*", r"\1. ", row.cells[0].text.strip())
            with self.subTest(item=heading):
                self.assertEqual(sections[heading], row.cells[1].text.strip())
        footer_section = reference.split("## B・C枠等の定型文", 1)[1].split("## 更新のしかた", 1)[0]
        footer_quote = "\n".join(line[2:] if line.startswith("> ") else ""
                                 for line in footer_section.splitlines() if line.startswith(">"))
        self.assertEqual(footer_quote.strip(), document.tables[1].cell(0, 0).text.strip())

    def test_builder_clears_all_example_cells_and_preserves_template(self) -> None:
        spec = importlib.util.spec_from_file_location("template_builder", SKILL / "scripts/build_return_docx_from_template.py")
        builder = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(builder)
        original = TEMPLATE.read_bytes()
        source = Document(TEMPLATE)
        lines = ["研究課題名【テンプレート検証】", "", "| 項目 | 対応要否 | 内容 |", "|---|---|---|"]
        for index, row in enumerate(source.tables[0].rows):
            label = re.sub(r"\s+", " ", row.cells[0].text).strip()
            status, text = ("要対応", "解析終了までの期間をご確認ください。") if index == 17 else ("対応不要", "")
            lines.append(f"| {label} | {status} | {text} |")
        lines.extend(["", "A．倫理審査申請システム質疑事項", "B．修正が必要な添付書類の再アップロード", "なし", "C．軽微な文言修正等を行った添付書類の再アップロード", "なし"])
        with tempfile.TemporaryDirectory() as tmp:
            draft = Path(tmp) / "04.md"
            draft.write_text("\n".join(lines), encoding="utf-8")
            output = Path(tmp) / "99_プレチェック出力_test" / "★倫理申請_修正対応依頼_テンプレート検証_1020.docx"
            builder.build_docx(TEMPLATE, draft, output)
            result = Document(output)
            for index, row in enumerate(result.tables[0].rows):
                expected = "解析終了までの期間をご確認ください。" if index == 17 else ""
                self.assertEqual(row.cells[1].text, expected)
            self.assertEqual(TEMPLATE.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
